"""Prepare private supervisor correspondence; this script never sends a message."""
import json
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]


def main():
    v=json.loads((ROOT/'results/revision_v2/manuscript_values.json').read_text())
    text=f'''# Draft response to the professor — for author review, not sent

Thank you for the detailed comments on “Multimodal Deep Learning for Real-Time Glucose Prediction with Zero/Few-Shot On-Premises LLM Explainability: External Validation on an Independent Adolescent Cohort.” This response refers to the 16-page manuscript_SUBMISSION_draft_mdpi.pdf supplied for review. The revision retains that title, the GlucoSense objective, the nine main sections and the original Sections 7.1–7.5. DDATSHR remains the main adolescent-cohort study. The completed DiaData analysis extends the larger-scale validation work already proposed in Section 4.2 and the original future-work list.

The original findings have been corrected where the underlying artifacts did not substantiate them; the numerical results below come from the completed experiments. The title is retained, but the Abstract and Introduction distinguish retrained external-cohort evaluation from frozen-model validation, and intended real-time use from an evaluated live service.

## 1. Deep models and modality discrepancy

All three reduced-input neural variants have now been trained and evaluated alongside Ridge, RF, XGBoost, persistence and linear trend on exactly the same patient/timestamp test manifests. The corrected cohorts contain {v['dd_n']} DDATSHR participants ({v['dd_samples']} origins) and {v['dia_n']} non-DDATSHR DiaData participants ({v['dia_samples']} origins). Each outer fold uses a separate validation patient; preprocessing and tuning exclude the test patient. Three seeds, restored checkpoints, direct simplex-constrained ensemble weights, CGM-only ablations, matched history sensitivity and source-held-out tests are saved.

RF macro MAE at 30/60/90 minutes is {v['dd_rf_mae']} mg/dL on DDATSHR and {v['dia_rf_mae']} mg/dL on DiaData. {v['primary_summary']}

The revised Discussion (Sections 7.5 and 8) no longer attributes the historical OhioT1DM gap to data scarcity or absent modalities. The earlier six-modal checkpoints and reconstructed streams are unavailable, and its result cannot be reproduced from this repository. Comparing newly fitted two-modal models with an unverified historical number does not isolate a modality effect. The new TFT variant is explicitly labelled an encoder inspired by TFT, rather than a reproduction of the full original architecture. The manuscript is therefore framed as a new retrospective evaluation, not locked-model external validation of the previous system.

## 2. Filter false negatives and apparent numerical reliability

The previous claim of perfect numeric grounding has been withdrawn. Its cases had null SHAP, invented ±15 mg/dL intervals, a single held-out participant, and a numerical filter that could confuse units and count missing claims as successes. The historical generation runner is disabled.

The new cases use actual held-out RF forecasts and TreeSHAP values whose additivity has been checked. The local generator must preserve a declared typed contract. Missing fields, wrong horizons/units/values, duplicate keys, undeclared advice and invalid categories are rejected. Free prose is not declared verified. Rejected LLM outputs are counted as rejected even though a deterministic template supplies a fallback.

On the final patient-separated cases, zero-shot acceptance is {v['zero_accept']} and few-shot acceptance is {v['few_accept']}; the deterministic comparator achieves {v['template_accept']}. The separately designated synthetic challenge families yielded {v['challenge_rejected']}/{v['challenge_total']} rejections. Crucially, all {v['boundary_total']} source-corruption cases were accepted when the incorrect source was faithfully copied. This gives an explicit example of end-to-end false reassurance outside the detector's contract. It cannot establish source truth, forecast accuracy, causal SHAP interpretation or clinical appropriateness. These limitations appear in Sections 3.3, 6.2–6.3 and 8.1. This is neither full FActScore nor clinician-rated safety.

## 3. Patient heterogeneity and participant 907

{v['heterogeneity_text']}

The new patient-profile table also contains recording duration, glucose and HR missingness, available demographics, rate of change and glycemic distribution. Associations are exploratory. We do not infer a physiological mechanism from a small cohort or equate ordinary Abbott 15-minute sampling with pathological fragmentation. Figures 4–5 and Section 7.3 present the corrected error distribution; Section 8 explains what can and cannot be attributed to patient physiology.

## 4. Grammar and equation indices

The original section sequence is retained, with corrected methods, result statements and grammar. Subject–verb agreement has been corrected, and the input notation consistently distinguishes participant p, forecast origin t, history index k and channel c. The prediction horizon h is explicit in each base forecast in Equation (2), and the RF-specific reference expectation is indexed by the held-out participant in Equation (3). Word equation indices are rendered as subscripts. The new source and generated Word document use the same result values.

## Additional changes needed to support these responses

Observed-only targets replace the interpolated DiaData targets. Upstream timestamp rounding is acknowledged with a conservative input delay. Device/clock exclusions and malformed rows are audited. Patient-level bootstrap intervals and paired tests replace inference over overlapping windows, with Holm correction across 18 primary tests. All Clarke zones are reported with the corrected lower-C boundary. New figures are supplied as vector files and 600-dpi PNGs. Related work now includes prior glucose/SHAP/LLM studies; broad novelty, clinical-safety, fine-tuning equivalence and zero-total-cost claims have been removed.

The earlier publication is explicitly disclosed. The revision does not reuse its artwork or present its unavailable result artifacts as verified. The clean manuscript, extensively highlighted review copy, result tables, code and verification record accompany this response. The highlight copy marks text differing from the user-supplied professor-reviewed PDF, not Word tracked deletions. REVISION_MAP_FOR_PROFESSOR.md explains corrections and displaced claims section by section.

## Additional strengthening while preserving the original paper

The separately labelled v3 work adds a new frozen DDATSHR-to-AZT1D glucose-only evaluation, a modern Chronos-2 comparator on unchanged primary test origins, and 276 RF sensitivity fits with three inner validation participants. AZT1D retains {v["azt_n"]} participants and {v["azt_samples"]} origins; one ambiguous CGM/BGM source field is excluded before scoring. This is not an adolescent-only cohort or a clinical trial. The models and stack were frozen before AZT1D scoring, and the primary 18-test family remains unchanged.

{v["additional_results_text"]}

A separate sequential replay measures actual prediction, TreeSHAP, local generation and validation/rendering on 30 fixed test cases. {v["latency_results_text"]} Tables 7–8 and Figure 7 document these additions within the original section structure. The original adolescent evaluation remains retraining-based; the new frozen glucose-only experiment does not authenticate the unavailable historical multimodal model.

## Before submission

Please review the complete revised research and confirm the author declarations, institutional secondary-data requirements and treatment of the related prior publication. This draft does not claim that you have already approved the revision, does not invent clinical validation, and has not been sent to any journal.
'''
    private=ROOT/'professor feedback';private.mkdir(exist_ok=True)
    (private/'RESPONSE_TO_PROFESSOR_REVISED.md').write_text(text)
    doc=Document()
    for paragraph in text.split('\n\n'):
        if paragraph.startswith('# '):doc.add_heading(paragraph[2:],0)
        elif paragraph.startswith('## '):doc.add_heading(paragraph[3:],1)
        else:doc.add_paragraph(paragraph)
    doc.save(private/'RESPONSE_TO_PROFESSOR_REVISED.docx')
    print('Private response drafts written; no correspondence sent.')

if __name__=='__main__':main()
